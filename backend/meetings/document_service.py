import os
from pypdf import PdfReader
import docx
import openpyxl

def format_file_size(num_bytes: int) -> str:
    """Formats bytes into human readable KB / MB."""
    for unit in ['B', 'KB', 'MB', 'GB']:
        if num_bytes < 1024.0:
            return f"{num_bytes:.1f} {unit}"
        num_bytes /= 1024.0
    return f"{num_bytes:.1f} GB"

def extract_text_from_file(file_path: str, ext: str) -> str:
    """
    Extracts plain text from PDF, DOCX, TXT, or XLSX files.
    """
    ext = ext.lower().strip()
    text = ""
    
    try:
        if ext == '.pdf':
            reader = PdfReader(file_path)
            pages_text = []
            for idx, page in enumerate(reader.pages):
                page_content = page.extract_text() or ''
                if page_content.strip():
                    pages_text.append(f"--- [Page {idx + 1}] ---\n{page_content.strip()}")
            text = "\n\n".join(pages_text)
            
        elif ext in ['.docx', '.doc']:
            doc = docx.Document(file_path)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            
            # Also extract tables from docx
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                    if row_text:
                        paragraphs.append(row_text)
                        
            text = "\n".join(paragraphs)
            
        elif ext in ['.xlsx', '.xls', '.csv']:
            wb = openpyxl.load_workbook(file_path, data_only=True)
            sheet_texts = []
            for sheet in wb.sheetnames:
                ws = wb[sheet]
                sheet_lines = [f"=== Sheet: {sheet} ==="]
                for row in ws.iter_rows(values_only=True):
                    row_vals = [str(v).strip() for v in row if v is not None and str(v).strip()]
                    if row_vals:
                        sheet_lines.append(" | ".join(row_vals))
                if len(sheet_lines) > 1:
                    sheet_texts.append("\n".join(sheet_lines))
            text = "\n\n".join(sheet_texts)
            
        elif ext in ['.txt', '.md', '.vtt', '.json']:
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                text = f.read()
                
    except Exception as e:
        print(f"[Document Text Extraction Error] ({file_path}): {e}")
        text = f"Error extracting document text: {str(e)}"
        
    return text.strip()
